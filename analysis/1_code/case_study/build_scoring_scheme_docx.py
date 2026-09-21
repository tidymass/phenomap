from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


PROJECT_ROOT = Path(__file__).resolve().parents[4]
OUT_DIR = PROJECT_ROOT / "analysis" / "3_result" / "case_study" / "cluster_5"
ZH_OUTPUT = OUT_DIR / "cluster_edge_biological_support_scoring_scheme.docx"
EN_OUTPUT = OUT_DIR / "cluster_edge_biological_support_scoring_scheme_en.docx"
FONT = "Arial Unicode MS"


def setup(doc):
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Pt(54)
    section.left_margin = section.right_margin = Pt(54)
    for style_name in ["Normal", "List Bullet"]:
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(10.5)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.1


def add_p(doc, text, size=10.5, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    return p


def add_h(doc, text):
    p = add_p(doc, text, 13, True)
    p.paragraph_format.space_before = Pt(10)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(10.5)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    doc.add_paragraph()


def build_zh():
    doc = Document()
    setup(doc)
    title = add_p(doc, "Cluster 5 Edge Biological Support Scoring Scheme", 18, True)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_p(doc, "目的：评价 Cluster 5 中由 metabolic module similarity 构建并绘制出的 phenotype-phenotype edge 是否具有已有生物学证据支持。")
    add_p(doc, "原则：评价对象是已有 cluster edge，不重新构建文献网络；文献只用于判断这些 edge 是否有 biological support。")
    add_p(doc, "Biological support 不要求两个节点共享完全相同的代谢机制。只要已有研究能够直接或间接说明两个节点之间存在合理生物学联系，即可记录为有一定支持。")

    add_h(doc, "一 节点范围")
    add_p(doc, "节点可包括 Disease、Biological process、Chemical 和 Exposure。每条 edge 按两个节点的类型进行文献检索和人工判定。")

    add_h(doc, "二 证据类型")
    add_table(doc, ["Evidence type", "定义"], [
        ["HUM", "人群、临床或流行病学研究直接同时涉及节点 A 和 B，并报告二者存在关联、共病、疾病风险、疾病进展或临床联系。"],
        ["EXP", "实验中直接操纵、给予或改变节点 A，并观察到节点 B 或与 B 相关的生物学变化。"],
        ["MECH", "已有研究明确描述 A 和 B 之间的分子、生化、代谢、调控、上下游或功能关系。"],
        ["SHARED", "没有研究直接连接 A 和 B，但 A、B 分别有文献支持与同一个较具体的 biological/metabolic mechanism 相关。"],
        ["STRUCT", "两个节点之间的关系主要来自 ontology、上下位、组成、同义、近义或定义关系，而不是独立 biological discovery。"],
    ])

    add_h(doc, "三 Direct evidence 判定")
    add_p(doc, "Direct 不要求证明因果关系。Direct 指已有研究本身同时涉及节点 A 和节点 B，并明确报告二者之间存在生物学、临床、流行病学或功能联系。")
    add_table(doc, ["Node pair", "可接受的 Direct relationship"], [
        ["Disease-Disease", "共病、风险关联、疾病进展、并发症、明确临床或流行病学关系。"],
        ["Disease-Exposure", "exposure 与疾病发生、风险、严重程度或进展存在关联。"],
        ["Disease-Chemical", "chemical 与疾病发生、治疗、毒性或疾病状态中的变化存在明确关系。"],
        ["Disease-Biological process", "process 明确参与疾病发生、发展或病理机制。"],
        ["Exposure-Chemical", "exposure 导致 chemical 水平或代谢发生变化。"],
        ["Exposure-Biological process", "exposure 明确诱导、抑制或改变该 process。"],
        ["Exposure-Exposure", "两种 exposure 存在明确交互、协同、拮抗或生物学相互作用。"],
        ["Chemical-Chemical", "存在代谢转化、precursor-product、酶促反应或直接生化关系。"],
        ["Chemical-Biological process", "chemical 明确参与、激活、抑制或调控该 process。"],
        ["Biological process-Biological process", "存在明确上下游、调控、功能耦联或 pathway interaction。"],
    ])
    add_p(doc, "只要存在 HUM、EXP 或 MECH 中至少一种明确 A-B 直接证据，即可将该 edge 归为 Direct。")

    add_h(doc, "四 Indirect evidence 判定")
    add_p(doc, "如果没有文献直接研究 A-B，但可以找到 A-M 以及 B-M，其中 M 是相对具体的 biological 或 metabolic mechanism，则认为该 edge 具有 Indirect biological support。")
    add_p(doc, "M 应具有机制特异性，通常应达到具体 pathway、metabolic route、molecular process、cellular process 或 functional program 的层级。")
    add_p(doc, "过于宽泛的机制单独出现时不足以定义为 Indirect，例如泛泛的 inflammation、oxidative stress、immune dysfunction、metabolic dysfunction、broad lipid metabolism 或 general mitochondrial dysfunction。")

    add_h(doc, "五 最终关系分类")
    add_table(doc, ["Final class", "定义"], [
        ["Direct", "已有研究直接同时涉及 A 和 B，并明确支持二者之间存在 biological relationship。证据可以来自 HUM、EXP 或 MECH。"],
        ["Indirect", "没有明确 A-B direct evidence，但 A、B 分别有可靠研究支持同一个较具体的 biological 或 metabolic mechanism。"],
        ["Limited", "能够找到一些可能的共同 biology，但联系过于宽泛或不够具体。"],
        ["No clear support", "经过检索后，没有发现明确 direct evidence，也没有发现可信 shared specific mechanism。该分类只表示未识别到明确先验证据，不代表二者一定不存在关系。"],
        ["Structural", "两个节点之间本身存在 parent-child ontology relationship、synonym/near-equivalent relationship、component relationship 或 definitional relationship。该类 edge 应单独标记。"],
    ])

    add_h(doc, "六 判定顺序")
    for item in [
        "首先判断两个节点是否属于 ontology、组成、同义、近义或定义关系。如果是，则归为 Structural。",
        "如果不是 Structural，则检索是否存在直接 A-B 文献。",
        "如果存在明确 HUM、EXP 或 MECH evidence，则归为 Direct。",
        "如果不存在 direct evidence，则分别检索 A 和 B 的 biological mechanisms。",
        "如果二者存在同一个较具体机制，则归为 Indirect。",
        "如果只有非常泛的共同机制，则归为 Limited。",
        "如果连可信的泛化联系也没有，则归为 No clear support。",
    ]:
        add_bullet(doc, item)
    add_p(doc, "判定优先级：Structural -> Direct -> Indirect -> Limited -> No clear support。")

    add_h(doc, "七 人工记录字段")
    for item in [
        "Node A, Type A, Node B, Type B",
        "PhenoMAP BMA score",
        "HUM, EXP, MECH, SHARED, STRUCT",
        "Final class",
        "Relation direction, if applicable",
        "Biological explanation",
        "Representative reference(s)",
        "Notes",
    ]:
        add_bullet(doc, item)

    add_h(doc, "八 结果汇总")
    add_p(doc, "对 Cluster 5 中所有待评价 edge 统计 Direct、Indirect、Limited、No clear support 和 Structural 的数量及比例。")
    add_p(doc, "Structural edges 应单独报告，不作为 ordinary biological validation positive edges。")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(ZH_OUTPUT)


def build_en():
    doc = Document()
    setup(doc)
    title = add_p(doc, "Cluster 5 Edge Biological Support Scoring Scheme", 18, True)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_p(doc, "Purpose: to evaluate whether phenotype-phenotype edges drawn from metabolic module similarity in Cluster 5 are supported by prior biological literature.")
    add_p(doc, "Principle: the evaluated unit is the existing cluster edge. Literature is used to assess support for these edges, not to reconstruct a new literature network.")
    add_p(doc, "Biological support does not require the two nodes to share exactly the same metabolic mechanism. Direct or indirect evidence for a plausible biological relationship is sufficient.")

    add_h(doc, "1 Node Scope")
    add_p(doc, "Nodes may include Disease, Biological process, Chemical, and Exposure. Each edge should be evaluated according to the types of its two nodes.")

    add_h(doc, "2 Evidence Types")
    add_table(doc, ["Evidence type", "Definition"], [
        ["HUM", "Human, clinical, or epidemiological evidence directly involves both node A and node B and reports an association, comorbidity, risk, progression, or clinical relationship."],
        ["EXP", "Experimental evidence manipulates, administers, or changes node A and observes node B or a biological change related to node B."],
        ["MECH", "Prior work describes a molecular, biochemical, metabolic, regulatory, upstream-downstream, or functional relationship between A and B."],
        ["SHARED", "No direct A-B study is found, but A and B are each supported by literature as related to the same specific biological or metabolic mechanism."],
        ["STRUCT", "The relationship mainly reflects ontology, hierarchy, composition, synonymy, near-equivalence, or definition rather than an independent biological discovery."],
    ])

    add_h(doc, "3 Direct Evidence")
    add_p(doc, "Direct evidence does not require proven causality. It means that a study directly involves both node A and node B and reports a biological, clinical, epidemiological, or functional relationship.")
    add_table(doc, ["Node pair", "Acceptable direct relationship"], [
        ["Disease-Disease", "Comorbidity, risk association, disease progression, complication, or clear clinical/epidemiological relationship."],
        ["Disease-Exposure", "Exposure is associated with disease onset, risk, severity, or progression."],
        ["Disease-Chemical", "Chemical is related to disease occurrence, treatment, toxicity, or disease-state change."],
        ["Disease-Biological process", "Process is involved in disease onset, progression, or pathology."],
        ["Exposure-Chemical", "Exposure changes chemical level or metabolism."],
        ["Exposure-Biological process", "Exposure induces, suppresses, or alters the process."],
        ["Exposure-Exposure", "Two exposures show interaction, synergy, antagonism, or biological interaction."],
        ["Chemical-Chemical", "Metabolic conversion, precursor-product relationship, enzymatic reaction, or direct biochemical relationship."],
        ["Chemical-Biological process", "Chemical participates in, activates, suppresses, or regulates the process."],
        ["Biological process-Biological process", "Upstream-downstream relationship, regulation, functional coupling, or pathway interaction."],
    ])
    add_p(doc, "If at least one clear HUM, EXP, or MECH source supports A-B directly, classify the edge as Direct.")

    add_h(doc, "4 Indirect Evidence")
    add_p(doc, "If no direct A-B study is found, but A-M and B-M are both supported and M is a specific biological or metabolic mechanism, the edge has indirect biological support.")
    add_p(doc, "M should be specific enough to represent a pathway, metabolic route, molecular process, cellular process, or functional program.")
    add_p(doc, "Very broad mechanisms alone are insufficient for Indirect classification, including generic inflammation, oxidative stress, immune dysfunction, metabolic dysfunction, broad lipid metabolism, or general mitochondrial dysfunction.")

    add_h(doc, "5 Final Classes")
    add_table(doc, ["Final class", "Definition"], [
        ["Direct", "A study directly involves both A and B and supports a biological relationship through HUM, EXP, or MECH evidence."],
        ["Indirect", "No clear A-B direct evidence is found, but A and B are each supported by reliable studies as related to the same specific biological or metabolic mechanism."],
        ["Limited", "Some shared biology is found, but the link is broad or nonspecific."],
        ["No clear support", "No clear direct evidence or credible shared specific mechanism is identified. This means no clear prior evidence was found, not that no relationship exists."],
        ["Structural", "The relationship reflects ontology, hierarchy, synonymy, near-equivalence, composition, or definition. These edges should be marked separately."],
    ])

    add_h(doc, "6 Scoring Order")
    for item in [
        "First check whether the two nodes have an ontological, compositional, synonymous, near-equivalent, or definitional relationship. If yes, classify as Structural.",
        "If not Structural, search for direct A-B literature.",
        "If clear HUM, EXP, or MECH evidence exists, classify as Direct.",
        "If direct evidence is absent, search mechanisms associated with A and B separately.",
        "If both nodes share the same specific mechanism, classify as Indirect.",
        "If only broad shared biology is found, classify as Limited.",
        "If no credible support is found, classify as No clear support.",
    ]:
        add_bullet(doc, item)
    add_p(doc, "Priority order: Structural -> Direct -> Indirect -> Limited -> No clear support.")

    add_h(doc, "7 Manual Annotation Fields")
    for item in [
        "Node A, Type A, Node B, Type B",
        "PhenoMAP BMA score",
        "HUM, EXP, MECH, SHARED, STRUCT",
        "Final class",
        "Relation direction, if applicable",
        "Biological explanation",
        "Representative reference(s)",
        "Notes",
    ]:
        add_bullet(doc, item)

    add_h(doc, "8 Summary")
    add_p(doc, "For all evaluated Cluster 5 edges, report the number and proportion of Direct, Indirect, Limited, No clear support, and Structural edges.")
    add_p(doc, "Structural edges should be reported separately and should not be interpreted as ordinary biological validation positive edges.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(EN_OUTPUT)


if __name__ == "__main__":
    build_zh()
    build_en()
    print(ZH_OUTPUT)
    print(EN_OUTPUT)
